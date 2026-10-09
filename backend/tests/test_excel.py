import os
import sqlite3
import pytest
from app.models.project import ProjectCreate
import tempfile
import openpyxl

def test_excel_ingestion(test_client, monkeypatch):
    # test_client sets up DB
    fd, temp_excel = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VideoIdeas"
    headers = ["id", "video_idea", "target_duration_seconds", "aspect_ratio", "status", "project_id"]
    ws.append(headers)
    ws.append(["EX-01", "Cat riding a roomba", 10, "16:9", "QUEUED", ""])
    ws.append(["EX-02", "Dog playing piano", 5, "9:16", "COMPLETED", ""])
    wb.save(temp_excel)
    
    # Override settings
    from app.core.config import settings
    monkeypatch.setattr(settings, "EXCEL_WORKBOOK_PATH", temp_excel)
    
    # Call ingest
    from app.services.excel_service import excel_service
    excel_service.read_and_ingest()
    
    from app.db.database import get_db
    with get_db() as db:
        projects = db.execute("SELECT * FROM projects WHERE excel_id = 'EX-01'").fetchall()
        assert len(projects) == 1
        assert projects[0]["video_idea"] == "Cat riding a roomba"
    
    # Test duplicate ingestion (idempotent)
    excel_service.read_and_ingest()
    with get_db() as db:
        projects = db.execute("SELECT * FROM projects WHERE excel_id = 'EX-01'").fetchall()
        assert len(projects) == 1 # Still 1
        
    # Update status
    excel_service.update_excel_status("EX-01", "COMPLETED")
    
    # Read back to verify
    wb_read = openpyxl.load_workbook(temp_excel)
    ws_read = wb_read["VideoIdeas"]
    assert ws_read.cell(row=2, column=5).value == "COMPLETED"
    
    # Negative test: invalid row
    ws_read.append(["EX-03", "", "invalid", "16:9", "QUEUED", ""]) # Invalid duration, empty idea
    wb_read.save(temp_excel)
    excel_service.read_and_ingest()
    wb_read2 = openpyxl.load_workbook(temp_excel)
    ws_read2 = wb_read2["VideoIdeas"]
    assert ws_read2.cell(row=4, column=5).value == "VALIDATION_FAILED"
    
    os.remove(temp_excel)
