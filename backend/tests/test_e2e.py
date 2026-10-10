import os
import sqlite3
import pytest
from fastapi.testclient import TestClient
import openpyxl

from app.main import app
from app.services.excel_service import excel_service

def test_full_e2e_workflow(test_client, monkeypatch):
    import subprocess
    from app.services.orchestration_service import orchestration_service
    from app.services.video_provider import video_provider
    from app.core.config import settings
    
    try:
        subprocess.run(["ffmpeg", "-version"], check=True, capture_output=True)
    except (FileNotFoundError, subprocess.CalledProcessError):
        pytest.skip("FFmpeg is not installed or not in PATH. Skipping genuine E2E test.")
    
    # 1. Setup mock excel
    temp_excel = os.path.join(settings.DATA_DIR, "video_ideas.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VideoIdeas"
    headers = ["id", "video_idea", "target_duration_seconds", "aspect_ratio", "status", "project_id"]
    ws.append(headers)
    ws.append(["EX-E2E", "A fully end-to-end test", 6, "16:9", "QUEUED", ""])
    wb.save(temp_excel)
    
    monkeypatch.setattr(settings, "EXCEL_WORKBOOK_PATH", str(temp_excel))
    
    # Setup deterministic video provider that runs real ffmpeg
    def mock_generate_video(project_id, shot_id, prompt, duration, aspect_ratio):
        shots_dir = os.path.join(settings.DATA_DIR, "projects", project_id, "shots")
        os.makedirs(shots_dir, exist_ok=True)
        out_file = os.path.join(shots_dir, f"{shot_id}.mp4")
        
        # Real ffmpeg call for a deterministic synthetic clip
        subprocess.run([
            "ffmpeg", "-y", "-f", "lavfi", "-i", f"color=c=blue:s=1920x1080:d={duration}", 
            "-c:v", "libx264", out_file
        ], check=True, capture_output=True)
        
        from app.db.database import get_db
        from app.models.project import ShotStatus
        with get_db() as db:
            db.execute("UPDATE shots SET status = ?, video_url = ? WHERE id = ?", 
                       (ShotStatus.COMPLETED.value, out_file, shot_id))
            db.commit()
            
        return out_file
        
    monkeypatch.setattr(video_provider, "generate_video_sync", mock_generate_video)

    triggered_projects = []
    def mock_trigger(pid):
        triggered_projects.append(pid)
        
    monkeypatch.setattr(orchestration_service, "trigger_pipeline", mock_trigger)
    
    # Ingest
    excel_service.read_and_ingest()
    
    # Verify it queued
    assert len(triggered_projects) == 1
    project_id = triggered_projects[0]
    
    # Run the pipeline synchronously to completion
    orchestration_service.run_pipeline_sync(project_id)
    
    # Verify final status
    from app.db.database import get_db
    with get_db() as db:
        proj = db.execute("SELECT status, final_video_url, sync_status FROM projects WHERE id = ?", (project_id,)).fetchone()
        assert proj["status"] == "READY_FOR_REVIEW"
        assert proj["final_video_url"] is not None
        assert proj["final_video_url"].endswith(".mp4")
        assert os.path.exists(proj["final_video_url"])
        assert proj["sync_status"] == "SUCCESS"
    
    # Check excel writeback
    wb_read = openpyxl.load_workbook(temp_excel)
    ws_read = wb_read["VideoIdeas"]
    assert ws_read.cell(row=2, column=5).value == "READY_FOR_REVIEW"

    # Now approve it
    response = test_client.post(f"/api/projects/{project_id}/approve")
    assert response.status_code == 200
    
    # Check it's COMPLETED in DB and Excel
    with get_db() as db:
        proj = db.execute("SELECT status FROM projects WHERE id = ?", (project_id,)).fetchone()
        assert proj["status"] == "COMPLETED"
        
    wb_read = openpyxl.load_workbook(temp_excel)
    ws_read = wb_read["VideoIdeas"]
    assert ws_read.cell(row=2, column=5).value == "COMPLETED"

def test_atomic_claim_concurrency(test_client, monkeypatch):
    """
    Tests that if two threads attempt to run the pipeline for the same project, 
    only one acquires the claim.
    """
    import threading
    from app.services.orchestration_service import orchestration_service
    from app.db.database import get_db
    import uuid
    
    project_id = str(uuid.uuid4())
    with get_db() as db:
        db.execute(
            "INSERT INTO projects (id, video_idea, target_duration, aspect_ratio, status) VALUES (?, ?, ?, ?, ?)",
            (project_id, "Concurrency test", 10, "16:9", "QUEUED")
        )
        db.commit()
        
    execution_count = [0]
    lock = threading.Lock()
    
    # We patch run_pipeline_sync just at the point where it does actual work,
    # or we can mock generate_storyboard. Let's mock generate_storyboard.
    from app.services.storyboard_service import storyboard_service
    original_generate = storyboard_service.generate_storyboard
    
    def mock_generate(*args, **kwargs):
        with lock:
            execution_count[0] += 1
        return original_generate(*args, **kwargs)
        
    monkeypatch.setattr(storyboard_service, "generate_storyboard", mock_generate)
    
    # Mock video_provider and assembly_service to avoid real FFmpeg
    from app.services.video_provider import video_provider
    from app.services.assembly_service import assembly_service
    monkeypatch.setattr(video_provider, "generate_video_sync", lambda *args, **kwargs: "/mock.mp4")
    monkeypatch.setattr(assembly_service, "assemble_shots", lambda *args, **kwargs: "/final.mp4")

    # Run 5 threads trying to execute the pipeline for the same project
    threads = []
    for _ in range(5):
        t = threading.Thread(target=orchestration_service.run_pipeline_sync, args=(project_id,))
        threads.append(t)
        t.start()
        
    for t in threads:
        t.join()
        
    # Only 1 execution should have happened!
    assert execution_count[0] == 1
