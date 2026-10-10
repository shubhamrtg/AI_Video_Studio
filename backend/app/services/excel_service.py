import os
import openpyxl
from pydantic import ValidationError
from datetime import datetime
import threading
import uuid
from app.core.config import settings
from app.db.database import get_db
from app.models.project import ProjectCreate, ProjectStatus
import logging

logger = logging.getLogger(__name__)

class ExcelService:
    def __init__(self):
        self.sheet_name = "VideoIdeas"
        self._lock = threading.Lock()

    @property
    def filepath(self):
        return settings.EXCEL_WORKBOOK_PATH

    def read_and_ingest(self):
        """Reads QUEUED rows from Excel and ingests them into the DB."""
        ingested_project_ids = []
        with self._lock:
            if not os.path.exists(self.filepath):
                logger.error(f"Excel file not found at {self.filepath}")
                return

            try:
                wb_data = openpyxl.load_workbook(self.filepath, data_only=True)
                if self.sheet_name not in wb_data.sheetnames:
                    logger.error(f"Sheet {self.sheet_name} not found in workbook")
                    return
                ws_data = wb_data[self.sheet_name]
                
                # Load a separate writable workbook to preserve formulas
                wb_write = openpyxl.load_workbook(self.filepath)
                ws_write = wb_write[self.sheet_name]
                
                # Map column names to indices
                headers = {}
                for col_idx, cell in enumerate(ws_data[1], 1):
                    if cell.value:
                        headers[cell.value] = col_idx
                        
                required_cols = ["id", "video_idea", "target_duration_seconds", "aspect_ratio", "status"]
                for col in required_cols:
                    if col not in headers:
                        logger.error(f"Missing required column in Excel: {col}")
                        return

                rows_ingested = 0
                
                with get_db() as db:
                    for row_idx in range(2, ws_data.max_row + 1):
                        status_val = ws_data.cell(row=row_idx, column=headers["status"]).value
                        if status_val == "QUEUED":
                            excel_id = str(ws_data.cell(row=row_idx, column=headers["id"]).value)
                            
                            # Check if already exists
                            existing = db.execute("SELECT id FROM projects WHERE excel_id = ?", (excel_id,)).fetchone()
                            if existing:
                                continue
                                
                            try:
                                project_data = ProjectCreate(
                                    excel_id=excel_id,
                                    video_idea=str(ws_data.cell(row=row_idx, column=headers["video_idea"]).value),
                                    target_duration=int(ws_data.cell(row=row_idx, column=headers["target_duration_seconds"]).value or 30),
                                    aspect_ratio=str(ws_data.cell(row=row_idx, column=headers["aspect_ratio"]).value or "16:9"),
                                    audio_policy=str(ws_data.cell(row=row_idx, column=headers["audio_policy"]).value or "silent").lower() if "audio_policy" in headers else "silent",
                                    visual_style=ws_data.cell(row=row_idx, column=headers.get("visual_style", -1)).value if "visual_style" in headers else None,
                                    language=ws_data.cell(row=row_idx, column=headers.get("language", -1)).value if "language" in headers else "English",
                                    voice_style=ws_data.cell(row=row_idx, column=headers.get("voice_style", -1)).value if "voice_style" in headers else None,
                                    target_platform=ws_data.cell(row=row_idx, column=headers.get("target_platform", -1)).value if "target_platform" in headers else None,
                                    priority=ws_data.cell(row=row_idx, column=headers.get("priority", -1)).value if "priority" in headers else "Normal"
                                )
                                
                                project_id = str(uuid.uuid4())
                                
                                db.execute("""
                                    INSERT INTO projects (
                                        id, excel_id, video_idea, target_duration, aspect_ratio, audio_policy,
                                        visual_style, language, voice_style, target_platform,
                                        priority, status
                                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                                """, (
                                    project_id, project_data.excel_id, project_data.video_idea,
                                    project_data.target_duration, project_data.aspect_ratio.value, project_data.audio_policy.value,
                                    project_data.visual_style, project_data.language,
                                    project_data.voice_style, project_data.target_platform,
                                    project_data.priority, ProjectStatus.QUEUED.value
                                ))
                                
                                # Write back the internal project_id
                                if "project_id" in headers:
                                    ws_write.cell(row=row_idx, column=headers["project_id"]).value = project_id
                                    
                                db.commit()
                                rows_ingested += 1
                                ingested_project_ids.append(project_id)
                                
                            except (ValueError, ValidationError) as e:
                                logger.error(f"Validation error on row {row_idx}: {e}")
                                if "error_message" in headers:
                                    ws_write.cell(row=row_idx, column=headers["error_message"]).value = str(e)
                                ws_write.cell(row=row_idx, column=headers["status"]).value = "VALIDATION_FAILED"
                
                if rows_ingested > 0 or "VALIDATION_FAILED" in [ws_write.cell(row=i, column=headers["status"]).value for i in range(2, ws_data.max_row+1)]:
                    try:
                        temp_path = self.filepath + ".tmp"
                        wb_write.save(temp_path)
                        os.replace(temp_path, self.filepath)
                        logger.info(f"Ingested {rows_ingested} new ideas.")
                    except PermissionError:
                        logger.error("Could not save Excel. Close the file if open.")
                        if os.path.exists(temp_path):
                            try:
                                os.remove(temp_path)
                            except OSError:
                                pass
                        
            except Exception as e:
                logger.error(f"Failed to ingest Excel: {e}")

        # Trigger orchestration background jobs for successfully ingested projects OUTSIDE the lock
        if ingested_project_ids:
            from app.services.orchestration_service import orchestration_service
            for pid in ingested_project_ids:
                try:
                    orchestration_service.trigger_pipeline(pid)
                except RuntimeError as e:
                    logger.error(f"Failed to queue project {pid} from excel: {e}")
                    with get_db() as db:
                        db.execute("UPDATE projects SET status = ?, error = ? WHERE id = ?", (ProjectStatus.FAILED.value, str(e), pid))
                        db.commit()

    def update_excel_status(self, excel_id: str, status: str, output_path: str = None, error: str = None):
        """Updates the status of a specific row in the Excel sheet."""
        with self._lock:
            if not os.path.exists(self.filepath):
                raise FileNotFoundError(f"Workbook {self.filepath} not found.")

            try:
                wb = openpyxl.load_workbook(self.filepath)
                if self.sheet_name not in wb.sheetnames:
                    raise ValueError(f"Worksheet {self.sheet_name} not found.")
                ws = wb[self.sheet_name]
                
                headers = {cell.value: col_idx for col_idx, cell in enumerate(ws[1], 1) if cell.value}
                
                if "id" not in headers or "status" not in headers:
                    raise ValueError("Required columns 'id' and 'status' not found in worksheet.")
                    
                row_found = False
                for row_idx in range(2, ws.max_row + 1):
                    cell_id = str(ws.cell(row=row_idx, column=headers["id"]).value)
                    if cell_id == excel_id:
                        row_found = True
                        ws.cell(row=row_idx, column=headers["status"]).value = status
                        
                        if output_path and "output_path" in headers:
                            ws.cell(row=row_idx, column=headers["output_path"]).value = output_path
                            
                        if error and "error_message" in headers:
                            ws.cell(row=row_idx, column=headers["error_message"]).value = error
                            
                        if status == "COMPLETED" and "completed_at" in headers:
                            ws.cell(row=row_idx, column=headers["completed_at"]).value = datetime.now().isoformat()
                            
                        temp_path = self.filepath + ".tmp"
                        try:
                            wb.save(temp_path)
                            os.replace(temp_path, self.filepath)
                        except Exception as save_err:
                            if os.path.exists(temp_path):
                                try:
                                    os.remove(temp_path)
                                except OSError:
                                    pass
                            raise RuntimeError(f"Failed to save Excel file: {save_err}")
                        break
                
                if not row_found:
                    raise ValueError(f"Excel ID {excel_id} not found in worksheet.")
            except Exception as e:
                # Re-raise to ensure caller knows sync failed
                raise RuntimeError(f"Excel sync failed: {str(e)}")

excel_service = ExcelService()
