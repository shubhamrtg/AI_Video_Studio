import os
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill

# Add backend directory to sys.path so we can import app modules
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.core.config import settings

def create_excel_template():
    filepath = settings.EXCEL_WORKBOOK_PATH
    if os.path.exists(filepath):
        print(f"Excel workbook already exists at {filepath}")
        return

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VideoIdeas"
    
    headers = [
        "id", "video_idea", "target_duration_seconds", "aspect_ratio",
        "visual_style", "language", "voice_style", "status",
        "target_platform", "reference_image_paths", "notes", "priority",
        "output_path", "error_message", "project_id", "created_at", "completed_at"
    ]
    
    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        
    # Write example row
    example_row = [
        "001", 
        "A professional assembles a miniature Porsche 911 GT3 RS using precision tools. Only hands and forearms are visible.",
        60,
        "9:16",
        "Photorealistic cinematic macro",
        "English",
        "No narration; realistic workshop ASMR",
        "QUEUED",
        "Instagram Reels",
        "",
        "Use macro lens look",
        "High",
        "",
        "",
        "",
        "",
        ""
    ]
    
    for col_idx, value in enumerate(example_row, 1):
        ws.cell(row=2, column=col_idx, value=value)
        
    # Set column widths
    ws.column_dimensions['B'].width = 80
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['G'].width = 40
    ws.column_dimensions['H'].width = 15
    
    wb.save(filepath)
    print(f"Created Excel template at {filepath}")

if __name__ == "__main__":
    create_excel_template()
